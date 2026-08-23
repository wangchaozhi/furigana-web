from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from ..models import ExportDocxRequest, LayoutSettings, Segment


CSS_PX_TO_PT = 72 / 96
COLUMN_GAP_TWIPS = 540  # 36 CSS px at 96 dpi


def _set_run_fonts(run, east_asia: str = "Yu Gothic", size: float = 12) -> None:
    run.font.name = "Arial"
    run.font.size = Pt(size)
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Arial")
    r_fonts.set(qn("w:hAnsi"), "Arial")
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _text_element(text: str) -> OxmlElement:
    t = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" ") or "  " in text:
        t.set(qn("xml:space"), "preserve")
    t.text = text
    return t


def _run_xml(text: str, size_half_points: int, east_asia: str) -> OxmlElement:
    r = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    fonts.set(qn("w:eastAsia"), east_asia)
    r_pr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(size_half_points))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(size_half_points))
    r_pr.append(sz_cs)
    r.append(r_pr)
    r.append(_text_element(text))
    return r


def _append_ruby(paragraph, base: str, ruby_text: str, base_size: int, ruby_size: int, east_asia: str) -> None:
    outer_run = OxmlElement("w:r")
    ruby = OxmlElement("w:ruby")

    ruby_pr = OxmlElement("w:rubyPr")
    align = OxmlElement("w:rubyAlign")
    align.set(qn("w:val"), "center")
    ruby_pr.append(align)

    hps = OxmlElement("w:hps")
    hps.set(qn("w:val"), str(ruby_size))
    ruby_pr.append(hps)

    hps_raise = OxmlElement("w:hpsRaise")
    # Keep a small, predictable gap between the base text and its phonetic
    # guide. A value of zero places both baselines together and makes Word
    # render the two strings on top of each other.
    hps_raise.set(qn("w:val"), str(max(0, base_size - 2)))
    ruby_pr.append(hps_raise)

    hps_base = OxmlElement("w:hpsBaseText")
    hps_base.set(qn("w:val"), str(base_size))
    ruby_pr.append(hps_base)

    lid = OxmlElement("w:lid")
    lid.set(qn("w:val"), "ja-JP")
    ruby_pr.append(lid)
    ruby.append(ruby_pr)

    rt = OxmlElement("w:rt")
    rt.append(_run_xml(ruby_text, ruby_size, east_asia))
    ruby.append(rt)

    ruby_base = OxmlElement("w:rubyBase")
    ruby_base.append(_run_xml(base, base_size, east_asia))
    ruby.append(ruby_base)

    # w:ruby is run content in WordprocessingML. Appending it directly to the
    # paragraph produces invalid OOXML that Word repairs when opening, often
    # dropping text or reflowing the surrounding content.
    outer_run.append(ruby)
    paragraph._p.append(outer_run)


def _append_segment(paragraph, seg: Segment, settings: LayoutSettings, east_asia: str) -> None:
    font_size_pt = settings.font_size * CSS_PX_TO_PT
    base_size = round(font_size_pt * 2)
    ruby_size = max(8, round(base_size * settings.ruby_scale))
    if seg.ruby:
        _append_ruby(paragraph, seg.text, seg.ruby, base_size, ruby_size, east_asia)
    elif seg.text:
        run = paragraph.add_run(seg.text)
        _set_run_fonts(run, east_asia, font_size_pt)


def _configure_section_layout(section, *, columns_count: int, vertical: bool) -> None:
    """Apply columns/writing direction without creating invalid sectPr XML."""
    sect_pr = section._sectPr

    # The default python-docx template already contains w:cols. Reuse it so a
    # two-column export does not contain duplicate section properties.
    columns = sect_pr.find(qn("w:cols"))
    if columns is None:
        columns = OxmlElement("w:cols")
        sect_pr.insert_element_before(
            columns,
            "w:formProt",
            "w:vAlign",
            "w:noEndnote",
            "w:titlePg",
            "w:textDirection",
            "w:bidi",
            "w:rtlGutter",
            "w:docGrid",
            "w:printerSettings",
            "w:sectPrChange",
        )
    columns.set(qn("w:num"), "2" if columns_count == 2 and not vertical else "1")
    columns.set(qn("w:space"), str(COLUMN_GAP_TWIPS))
    if columns_count == 2 and not vertical:
        columns.set(qn("w:sep"), "1")
    else:
        columns.attrib.pop(qn("w:sep"), None)

    text_direction = sect_pr.find(qn("w:textDirection"))
    if vertical:
        if text_direction is None:
            text_direction = OxmlElement("w:textDirection")
            sect_pr.insert_element_before(
                text_direction,
                "w:bidi",
                "w:rtlGutter",
                "w:docGrid",
                "w:printerSettings",
                "w:sectPrChange",
            )
        text_direction.set(qn("w:val"), "tbRl")
    elif text_direction is not None:
        sect_pr.remove(text_direction)


def build_docx(request: ExportDocxRequest) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    margin_mm = request.layout.page_margin * 0.32
    section.top_margin = Mm(margin_mm)
    section.bottom_margin = Mm(margin_mm)
    section.left_margin = Mm(margin_mm)
    section.right_margin = Mm(margin_mm)
    font_names = {"gothic": "Yu Gothic", "mincho": "Yu Mincho", "system": "Arial"}
    east_asia = font_names[request.layout.font_family]

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(request.layout.font_size * CSS_PX_TO_PT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)

    has_title_block = bool(request.meta.title or request.meta.artist or request.meta.year)

    if request.meta.title:
        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(2)
        title.paragraph_format.keep_with_next = True
        run = title.add_run(request.meta.title)
        run.bold = True
        run.font.size = Pt(22)
        _set_run_fonts(run, east_asia, 22)
        run.font.size = Pt(22)

    subtitle_parts = []
    if request.meta.artist:
        subtitle_parts.append(f"Song by {request.meta.artist}")
    if request.meta.year:
        subtitle_parts.append(request.meta.year)
    if subtitle_parts:
        sub = document.add_paragraph(" · ".join(subtitle_parts))
        sub.paragraph_format.space_after = Pt(14)
        sub.paragraph_format.keep_with_next = True
        for run in sub.runs:
            _set_run_fonts(run, east_asia, 10.5)
            run.font.size = Pt(10.5)

    # Match the browser preview: the title block stays horizontal and
    # single-column while only the lyrics body uses columns/vertical writing.
    # Column changes can use a continuous break. Word-compatible renderers
    # require a page break when the section writing direction changes.
    body_section = section
    if has_title_block and (request.layout.columns == 2 or request.layout.vertical):
        _configure_section_layout(section, columns_count=1, vertical=False)
        section_start = WD_SECTION.NEW_PAGE if request.layout.vertical else WD_SECTION.CONTINUOUS
        body_section = document.add_section(section_start)
        if request.layout.vertical:
            section_type = body_section._sectPr.find(qn("w:type"))
            if section_type is None:
                section_type = OxmlElement("w:type")
                body_section._sectPr.insert_element_before(
                    section_type,
                    "w:pgSz",
                    "w:pgMar",
                    "w:paperSrc",
                    "w:pgBorders",
                    "w:lnNumType",
                    "w:pgNumType",
                    "w:cols",
                )
            section_type.set(qn("w:val"), "nextPage")
    _configure_section_layout(
        body_section,
        columns_count=request.layout.columns,
        vertical=request.layout.vertical,
    )

    for line in request.lines:
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0 if line.translation and request.translation_language != "none" else 2)
        body_size_pt = request.layout.font_size * CSS_PX_TO_PT
        p.paragraph_format.line_spacing = Pt(body_size_pt * request.layout.line_spacing)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        p.paragraph_format.keep_together = True
        if not line.segments:
            p.add_run("")
            continue
        for seg in line.segments:
            _append_segment(p, seg, request.layout, east_asia)
        if line.translation and request.translation_language != "none":
            p.paragraph_format.keep_with_next = True
            translation = document.add_paragraph()
            translation.paragraph_format.space_after = Pt(6)
            translation.paragraph_format.keep_together = True
            translation_size_pt = max(9, request.layout.font_size * 0.72 * CSS_PX_TO_PT)
            translation.paragraph_format.line_spacing = Pt(translation_size_pt * 1.45)
            translation.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            run = translation.add_run(line.translation)
            _set_run_fonts(run, east_asia, translation_size_pt)
            run.font.color.rgb = RGBColor(90, 90, 84)

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
